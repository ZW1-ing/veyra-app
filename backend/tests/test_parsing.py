"""文档解析与上传：PDF / DOCX / txt 都要能抽出文字并入库。"""

import io

from docx import Document as DocxDocument
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from app.kb.parsing import (
    DocumentParseError,
    UnsupportedDocument,
    detect_suffix,
    extract_text,
)


def make_pdf(text: str) -> bytes:
    # 默认的 Helvetica 不含中文字形，直接画中文会写进一堆乱码；
    # 用 reportlab 内置的 CID 字体，测试才覆盖真实场景（中文 PDF）
    if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.setFont("STSong-Light", 14)
    pdf.drawString(72, 720, text)
    pdf.save()
    return buffer.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        doc_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, value in enumerate(row):
                doc_table.cell(row_index, col_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_detect_suffix_handles_paths_and_case():
    assert detect_suffix("C:\\Users\\me\\报告.PDF") == ".pdf"
    assert detect_suffix("notes.md") == ".md"
    assert detect_suffix("no-extension") == ""


def test_extract_plain_text_handles_utf8_and_gbk():
    assert extract_text("a.txt", "中文内容".encode()) == "中文内容"
    # Windows 上导出的老文件常是 GBK
    assert extract_text("b.txt", "中文内容".encode("gb18030")) == "中文内容"


def test_extract_pdf_reads_real_pdf():
    text = extract_text("sample.pdf", make_pdf("Veyra PDF parsing works"))
    assert "Veyra PDF parsing works" in text


def test_extract_docx_reads_paragraphs_and_tables():
    content = make_docx(
        ["第一段：项目说明", "第二段：使用方式"],
        table=[["字段", "含义"], ["kb", "知识库"]],
    )
    text = extract_text("sample.docx", content)

    assert "第一段：项目说明" in text
    assert "第二段：使用方式" in text
    # 表格里的关键信息也要被抽出来
    assert "字段 | 含义" in text
    assert "kb | 知识库" in text


def test_unsupported_and_broken_files_raise_readable_errors():
    try:
        extract_text("photo.png", b"\x89PNG")
    except UnsupportedDocument as exc:
        assert "暂不支持" in str(exc)
    else:
        raise AssertionError("不支持的格式应当抛 UnsupportedDocument")

    try:
        extract_text("broken.pdf", b"not a pdf at all")
    except DocumentParseError as exc:
        assert "PDF" in str(exc)
    else:
        raise AssertionError("损坏的 PDF 应当抛 DocumentParseError")


def test_upload_endpoint_ingests_pdf(client):
    response = client.post(
        "/kb/documents/upload",
        files={"file": ("report.pdf", make_pdf("上传解析验证：离线优先"), "application/pdf")},
    )
    assert response.status_code == 201, response.text
    body = response.json()

    assert body["name"] == "report.pdf"
    assert body["source_type"] == "file"
    assert body["char_count"] > 0
    # 解析出来的正文进了库，检索应该能命中。
    # 用完整短语查询：短查询（比如「离线优先」四个字）在阈值附近不稳，容易误判成 bug。
    hits = client.post("/kb/search", json={"query": "上传解析验证 离线优先"}).json()
    assert any(hit["document_name"] == "report.pdf" for hit in hits)


def test_upload_endpoint_rejects_unsupported_type(client):
    response = client.post(
        "/kb/documents/upload",
        files={"file": ("a.png", b"\x89PNG", "image/png")},
    )
    assert response.status_code == 400
    assert "暂不支持" in response.json()["detail"]
